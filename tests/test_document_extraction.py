import io

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.services.document_service import extract_document_text


def test_extracts_chinese_plain_text_and_markdown():
    assert "蓝鲸计划" in extract_document_text(
        "蓝鲸计划总预算 4200 万元".encode(),
        "plan.txt",
        "text/plain",
    )
    assert "灰度比例 37%" in extract_document_text(
        "会议结论：灰度比例 37%".encode(),
        "meeting.md",
        "text/markdown",
    )


def test_extracts_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("文档正文包含项目代号：蓝鲸计划")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格预算：4200万元"
    stream = io.BytesIO()
    document.save(stream)

    text = extract_document_text(stream.getvalue(), "plan.docx")

    assert "蓝鲸计划" in text
    assert "4200万元" in text


def test_extracts_xlsx_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预测"
    sheet.append(["区域", "营收"])
    sheet.append(["东南区", 18300000])
    stream = io.BytesIO()
    workbook.save(stream)

    text = extract_document_text(stream.getvalue(), "forecast.xlsx")

    assert "工作表：预测" in text
    assert "东南区\t18300000" in text


def test_extracts_pptx_text():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "星图发布会"
    stream = io.BytesIO()
    presentation.save(stream)

    text = extract_document_text(stream.getvalue(), "deck.pptx")

    assert "星图发布会" in text


def test_unknown_binary_returns_empty_text():
    assert extract_document_text(b"\x00\x01\x02", "legacy.doc") == ""
