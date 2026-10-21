"""从常见文档中抽取可检索的正文。"""

import io
from pathlib import Path

from app.config import settings

_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
}


def extract_document_text(
    data: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:
    extension = Path(filename).suffix.lower()
    if extension in _TEXT_EXTENSIONS or (content_type or "").startswith("text/"):
        return _limit(_clean_text(_decode_text(data)))
    if extension == ".pdf" or content_type == "application/pdf":
        return _limit(_clean_text(_extract_pdf(data)))
    if extension == ".docx":
        return _limit(_clean_text(_extract_docx(data)))
    if extension == ".xlsx":
        return _limit(_clean_text(_extract_xlsx(data)))
    if extension == ".pptx":
        return _limit(_clean_text(_extract_pptx(data)))
    return ""


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    parts.append("\t".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"幻灯片 {index}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return "\n".join(parts)


def _clean_text(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    cleaned = "\n".join(lines).strip()
    while "\n\n\n\n" in cleaned:
        cleaned = cleaned.replace("\n\n\n\n", "\n\n\n")
    return cleaned


def _limit(text: str) -> str:
    return text[: settings.MAX_DOCUMENT_TEXT_CHARS].strip()
